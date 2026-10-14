import os
import uuid
import json
import logging
from threading import BoundedSemaphore
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone

from flask import Flask, request, jsonify, send_file, Response
from flask_cors import CORS

from utils.lms_exporter import (
    export_to_moodle_xml,
    export_to_qti_21,
    export_to_google_forms,
    export_to_docx,
)

from config import (
    FLASK_HOST,
    FLASK_PORT,
    FLASK_DEBUG,
    T5_MODEL_NAME,
    BERT_MODEL_NAME,
    TORCH_NUM_THREADS,
    MAX_CONCURRENT_GENERATIONS,
    CORS_ORIGINS,
)
from utils.db_backup import backup_database
from database.db import (
    get_connection,
    init_db,
    save_paper,
    get_all_papers,
    get_paper_by_id,
    delete_paper_by_id,
    update_paper,
    get_pyq_stats_by_subject,
    get_pyq_questions_paginated,
    get_all_pyq_analytics,
    get_solutions_for_questions,
    create_user,
    get_user_by_email,
    create_assessment,
    get_assessments_for_student,
    save_attempt,
    get_attempts_for_student,
    get_all_attempts,
    update_attempt_evaluation,
)
from utils.nlp_processor import NLPProcessor
from utils.ai_engine import AIEngine
from utils.smart_selector import SmartSelector
from utils.paper_structurer import PaperStructurer
from utils.pdf_generator import PDFGenerator
from utils.bloom_classifier import BloomClassifier
from utils.rag_engine import VectorRAGEngine
from utils.answer_evaluator import AnswerEvaluator
from utils.mcq_generator import MCQGenerator
from utils.llm_gateway import LLMGateway

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)
logger = logging.getLogger(__name__)

# Throttling & Concurrency controls to protect worker CPU
try:
    import torch
    torch.set_num_threads(TORCH_NUM_THREADS)
    logger.info("Configured PyTorch CPU thread pool: %d threads", TORCH_NUM_THREADS)
except Exception:
    pass

# Concurrency semaphore prevents parallel heavy ML generations from starving CPU / crashing OOM
generation_semaphore = BoundedSemaphore(value=MAX_CONCURRENT_GENERATIONS)
task_executor = ThreadPoolExecutor(max_workers=MAX_CONCURRENT_GENERATIONS)
GENERATION_TASKS = {}


# Initialize Flask app
app = Flask(__name__)
allowed_origins = (
    "*"
    if CORS_ORIGINS.strip() == "*"
    else [origin.strip() for origin in CORS_ORIGINS.split(",") if origin.strip()]
)
CORS(app, resources={r"/api/*": {"origins": allowed_origins}})  # type: ignore

# Initialize components
llm_gateway = LLMGateway()
nlp = NLPProcessor()
ai_engine = AIEngine(T5_MODEL_NAME, BERT_MODEL_NAME, llm_gateway=llm_gateway)
selector = SmartSelector(ai_engine)
structurer = PaperStructurer()
pdf_gen = PDFGenerator()
bloom = BloomClassifier(llm_gateway=llm_gateway)
rag_engine = VectorRAGEngine(ai_engine)
evaluator = AnswerEvaluator(ai_engine=ai_engine, llm_gateway=llm_gateway)
mcq_generator = MCQGenerator(llm_gateway=llm_gateway)

# Initialize database
init_db()

# Eagerly load AI models at startup so the first HTTP request doesn't
# trigger a multi-minute model download / load that times out the client.
logger.info("Pre-loading AI models at startup (this may take a few minutes on first run)...")
try:
    ai_engine.load_models()
    logger.info("AI models loaded successfully.")
    # Build Vector RAG index
    try:
        pyqs = get_pyq_questions_paginated(limit=3000)["questions"]
        rag_engine.build_index(pyqs)
    except Exception as _rag_err:
        logger.warning("Vector RAG indexing deferred: %s", _rag_err)
except Exception as _exc:
    logger.warning("Model pre-load encountered an error (fallback will be used): %s", _exc)


def _execute_paper_generation(data: dict, progress_cb=None) -> dict:
    subject = data.get("subject", "").strip()
    syllabus = data.get("syllabus", "").strip()
    exam_pattern = data.get("exam_pattern", "standard")
    total_marks = int(data.get("total_marks", 80))
    duration_minutes = int(data.get("duration_minutes", 180))
    num_questions = int(data.get("num_questions", 9))
    difficulty_distribution = data.get("difficulty_distribution", {"easy": 30, "medium": 50, "hard": 20})
    organization_name = data.get("organization_name", "").strip()
    semester = data.get("semester", "").strip()

    short_questions_count = int(data.get("short_questions_count", 5))
    short_questions_marks = int(data.get("short_questions_marks", 2))
    short_questions_total = int(data.get("short_questions_total", 10))
    short_questions_choice_generate = int(data.get("short_questions_choice_generate", 7))
    short_questions_choice_attempt = int(data.get("short_questions_choice_attempt", 5))
    long_questions_count = int(data.get("long_questions_count", 8))
    long_questions_marks = int(data.get("long_questions_marks", 15))
    long_questions_total = int(data.get("long_questions_total", 60))

    if not subject or not syllabus:
        raise ValueError("Subject and syllabus are required")
    if len(syllabus) < 10:
        raise ValueError("Syllabus must be at least 10 characters")

    logger.info("Generating paper for subject: %s, pattern: %s", subject, exam_pattern)

    # Step 1: NLP — Extract topics and units
    if progress_cb:
        progress_cb(15, "Analyzing syllabus & extracting unit topics...")
    topics = nlp.extract_topics(syllabus)
    units = nlp.extract_units(syllabus)
    unit_topic_map = nlp.map_topics_to_units(syllabus)
    important_topics = nlp.get_important_topics(syllabus, top_n=20)

    # Step 2: AI — Generate questions using Hybrid RAG
    if progress_cb:
        progress_cb(40, "Retrieving grounded PYQs via Hybrid RAG & synthesizing questions...")
    all_questions = []

    if exam_pattern == "certification":
        short_questions_per_topic = max(1, (short_questions_count * 2) // max(len(important_topics), 1))
        long_questions_per_topic = max(1, (long_questions_count * 2) // max(len(important_topics), 1))

        for topic in important_topics:
            topic_unit = "Unit 1"
            for unit_name, unit_topics in unit_topic_map.items():
                unit_topics_list: list[str] = list(unit_topics)
                if any(topic.lower() in t.lower() or t.lower() in topic.lower() for t in unit_topics_list):
                    topic_unit = unit_name
                    break

            grounding_ctx = rag_engine.get_grounding_context(topic, subject=subject, top_k=2)
            try:
                pyq_short = ai_engine.generate_questions_with_rag_pyq(
                    subject, f"{topic} (short answer)", short_questions_per_topic,
                    grounding_context=grounding_ctx
                )
            except Exception:
                pyq_short = ai_engine._generate_fallback_questions_dict(topic, short_questions_per_topic, short_mode=True)

            for q_data in pyq_short:
                all_questions.append({
                    "id": str(uuid.uuid4()),
                    "text": q_data["text"],
                    "marks": short_questions_marks,
                    "difficulty": q_data["difficulty"],
                    "unit": topic_unit,
                    "topic": topic,
                    "question_type": "short",
                    "source": q_data.get("source", "Hybrid RAG Engine"),
                })

            try:
                pyq_long = ai_engine.generate_questions_with_rag_pyq(
                    subject, f"{topic} (long answer)", long_questions_per_topic,
                    grounding_context=grounding_ctx
                )
            except Exception:
                pyq_long = ai_engine._generate_fallback_questions_dict(topic, long_questions_per_topic, short_mode=False)

            for q_data in pyq_long:
                all_questions.append({
                    "id": str(uuid.uuid4()),
                    "text": q_data["text"],
                    "marks": long_questions_marks,
                    "difficulty": q_data["difficulty"],
                    "unit": topic_unit,
                    "topic": topic,
                    "question_type": "long",
                    "source": q_data.get("source", "Hybrid RAG Engine"),
                })
    else:
        questions_per_topic = max(1, (num_questions * 2) // max(len(important_topics), 1))
        for topic in important_topics:
            topic_unit = "Unit 1"
            for unit_name, unit_topics in unit_topic_map.items():
                unit_topics_list_gen: list[str] = list(unit_topics)
                if any(topic.lower() in t.lower() or t.lower() in topic.lower() for t in unit_topics_list_gen):
                    topic_unit = unit_name
                    break

            grounding_ctx = rag_engine.get_grounding_context(topic, subject=subject, top_k=2)
            try:
                pyq_generated = ai_engine.generate_questions_with_rag_pyq(
                    subject, topic, questions_per_topic, grounding_context=grounding_ctx
                )
            except Exception:
                pyq_generated = ai_engine._generate_fallback_questions_dict(topic, questions_per_topic)

            for q_data in pyq_generated:
                all_questions.append({
                    "id": str(uuid.uuid4()),
                    "text": q_data["text"],
                    "marks": q_data["marks"],
                    "difficulty": q_data["difficulty"],
                    "unit": topic_unit,
                    "topic": topic,
                    "question_type": q_data.get("question_type") or "descriptive",
                    "source": q_data.get("source", "Hybrid RAG Engine"),
                })

    # Step 3: Selection
    if progress_cb:
        progress_cb(70, "Eliminating duplicates & selecting balanced questions...")
    if exam_pattern == "certification":
        short_questions = [q for q in all_questions if q.get("question_type") == "short"]
        long_questions = [q for q in all_questions if q.get("question_type") == "long"]
        selected_short = short_questions[:short_questions_choice_generate]
        selected_long = long_questions[:long_questions_count]
        selected = selected_short + selected_long
    else:
        try:
            selected = selector.select_questions(
                all_questions, difficulty_distribution, num_questions
            )
        except Exception as e:
            logger.warning("Smart selection fallback: %s", e)
            import itertools
            selected = list(itertools.islice(all_questions, num_questions))

    # Step 4: Structure Paper
    if progress_cb:
        progress_cb(85, "Structuring paper sections and marks...")
    try:
        if exam_pattern == "certification":
            sections = structurer.structure_certification_paper(
                selected,
                short_questions_count, short_questions_marks, short_questions_total,
                short_questions_choice_generate, short_questions_choice_attempt,
                long_questions_count, long_questions_marks, long_questions_total
            )
        else:
            sections = structurer.structure_paper(selected, exam_pattern, total_marks)
    except Exception as e:
        logger.warning("Paper structuring fallback: %s", e)
        sections = [{
            "name": "Section A",
            "instructions": "Answer all questions",
            "questions": selected,
            "total_marks": sum(q["marks"] for q in selected)
        }]

    # Step 5: Bloom's Taxonomy
    if progress_cb:
        progress_cb(95, "Tagging cognitive depth with Bloom's taxonomy...")
    selected = bloom.tag_questions(selected)
    for sec in sections:
        sec_questions = sec.get("questions")
        if isinstance(sec_questions, list):
            sec["questions"] = bloom.tag_questions(sec_questions)

    paper_id = str(uuid.uuid4())
    paper = {
        "id": paper_id,
        "subject": subject,
        "organization_name": organization_name,
        "semester": semester,
        "syllabus": syllabus,
        "exam_pattern": exam_pattern,
        "total_marks": total_marks,
        "duration_minutes": duration_minutes,
        "num_questions": len(selected),
        "difficulty_distribution": difficulty_distribution,
        "questions": selected,
        "sections": sections,
        "syllabus_topics": important_topics,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }

    save_paper(paper)
    if progress_cb:
        progress_cb(100, "Paper generation complete!")
    return paper


@app.route("/api/generate", methods=["POST"])
def generate_paper():
    acquired = generation_semaphore.acquire(blocking=True, timeout=60)
    if not acquired:
        logger.warning("Generation request rejected: server at capacity")
        return jsonify({
            "error": "The generation engine is currently operating at peak capacity. Please retry shortly."
        }), 429

    try:
        data = request.get_json()
        if not data:
            return jsonify({"error": "No JSON data provided"}), 400
        paper = _execute_paper_generation(data)
        return jsonify(paper)
    except ValueError as val_err:
        return jsonify({"error": str(val_err)}), 400
    except Exception as e:
        logger.exception("Error generating paper")
        return jsonify({"error": str(e)}), 500
    finally:
        generation_semaphore.release()


@app.route("/api/generate/async", methods=["POST"])
def generate_paper_async():
    """Submit a paper generation job asynchronously without HTTP timeout risk."""
    data = request.get_json()
    if not data:
        return jsonify({"error": "No JSON data provided"}), 400

    subject = data.get("subject", "").strip()
    syllabus = data.get("syllabus", "").strip()
    if not subject or not syllabus:
        return jsonify({"error": "Subject and syllabus are required"}), 400

    task_id = str(uuid.uuid4())
    now_iso = datetime.now(timezone.utc).isoformat()
    GENERATION_TASKS[task_id] = {
        "task_id": task_id,
        "status": "queued",
        "progress": 5,
        "stage": "Job queued in background worker pool...",
        "result": None,
        "error": None,
        "created_at": now_iso,
        "updated_at": now_iso,
    }

    def _async_worker():
        acquired = generation_semaphore.acquire(blocking=True, timeout=120)
        if not acquired:
            GENERATION_TASKS[task_id]["status"] = "failed"
            GENERATION_TASKS[task_id]["error"] = "Server worker capacity exhausted."
            return

        try:
            def _update_progress(percent: int, stage_desc: str):
                GENERATION_TASKS[task_id]["progress"] = percent
                GENERATION_TASKS[task_id]["stage"] = stage_desc
                GENERATION_TASKS[task_id]["status"] = "processing"
                GENERATION_TASKS[task_id]["updated_at"] = datetime.now(timezone.utc).isoformat()

            paper_result = _execute_paper_generation(data, progress_cb=_update_progress)
            GENERATION_TASKS[task_id]["status"] = "completed"
            GENERATION_TASKS[task_id]["progress"] = 100
            GENERATION_TASKS[task_id]["stage"] = "Examination paper created successfully!"
            GENERATION_TASKS[task_id]["result"] = paper_result
            GENERATION_TASKS[task_id]["updated_at"] = datetime.now(timezone.utc).isoformat()
        except Exception as async_err:
            logger.exception("Async paper generation failed")
            GENERATION_TASKS[task_id]["status"] = "failed"
            GENERATION_TASKS[task_id]["error"] = str(async_err)
            GENERATION_TASKS[task_id]["updated_at"] = datetime.now(timezone.utc).isoformat()
        finally:
            generation_semaphore.release()

    task_executor.submit(_async_worker)
    return jsonify({
        "task_id": task_id,
        "status": "queued",
        "poll_url": f"/api/tasks/{task_id}"
    }), 202


@app.route("/api/tasks/<task_id>", methods=["GET"])
def get_task_status(task_id: str):
    """Retrieve async generation task status and result."""
    task = GENERATION_TASKS.get(task_id)
    if not task:
        return jsonify({"error": "Task not found"}), 404
    return jsonify(task)



@app.route("/api/papers", methods=["GET"])
def list_papers():
    try:
        papers = get_all_papers()
        return jsonify(papers)
    except Exception as e:
        logger.exception("Error listing papers")
        return jsonify({"error": str(e)}), 500


@app.route("/api/auth/signup", methods=["POST"])
def signup():
    data = request.get_json() or {}
    name = str(data.get("name", "")).strip()
    email = str(data.get("email", "")).strip().lower()
    password = str(data.get("password", ""))
    role = data.get("role")
    if not name or "@" not in email or len(password) < 6 or role not in ("student", "teacher"):
        return jsonify({"error": "Name, valid email, role, and a 6-character password are required"}), 400
    try:
        user = create_user(name, email, password, role, data.get("enrollments", []))
        return jsonify(user), 201
    except Exception as exc:
        if "UNIQUE constraint failed" in str(exc):
            return jsonify({"error": "An account with this email already exists"}), 409
        logger.exception("Error creating user")
        return jsonify({"error": str(exc)}), 500


@app.route("/api/auth/login", methods=["POST"])
def login():
    data = request.get_json() or {}
    user = get_user_by_email(str(data.get("email", "")), str(data.get("password", "")))
    if not user or user["role"] != data.get("role"):
        return jsonify({"error": "Invalid email, password, or role"}), 401
    return jsonify(user)


@app.route("/api/assessments", methods=["POST"])
def publish_assessment():
    data = request.get_json() or {}
    required = ("id", "title", "paper_id", "course", "subject", "subject_name", "session_code", "created_at")
    if any(not data.get(key) for key in required):
        return jsonify({"error": "Assessment details are incomplete"}), 400
    try:
        return jsonify(create_assessment(data)), 201
    except Exception as exc:
        logger.exception("Error publishing assessment")
        return jsonify({"error": str(exc)}), 500


@app.route("/api/student/assessments", methods=["GET"])
def student_assessments():
    email = request.args.get("email", "").strip().lower()
    user = get_user_by_email(email)
    if not user or user["role"] != "student":
        return jsonify({"error": "Student account not found"}), 404
    return jsonify(get_assessments_for_student(user["id"]))


@app.route("/api/student/attempts", methods=["GET", "POST"])
def student_attempts():
    if request.method == "GET":
        user = get_user_by_email(request.args.get("email", ""))
        if not user or user["role"] != "student":
            return jsonify({"error": "Student account not found"}), 404
        return jsonify(get_attempts_for_student(user["id"]))

    data = request.get_json() or {}
    user = get_user_by_email(str(data.get("student_email", "")))
    if not user or user["role"] != "student":
        return jsonify({"error": "Student account not found"}), 404
    assessment_ids = {item["id"] for item in get_assessments_for_student(user["id"])}
    if data.get("assessment_id") not in assessment_ids:
        return jsonify({"error": "You are not enrolled in this assessment"}), 403
    existing_attempts = get_attempts_for_student(user["id"])
    if any(attempt["assessment_id"] == data.get("assessment_id") for attempt in existing_attempts):
        return jsonify({"error": "You have already submitted this assessment. Each test can only be attempted once."}), 409
    data["student_id"] = user["id"]
    return jsonify(save_attempt(data)), 201


@app.route("/api/teacher/attempts", methods=["GET"])
def teacher_attempts():
    return jsonify(get_all_attempts())


@app.route("/api/teacher/analytics", methods=["GET"])
def teacher_learning_analytics():
    """Aggregate real per-question AI grades into concept mastery insights."""
    attempts = get_all_attempts()
    concepts = {}
    evaluated_attempts = 0
    for attempt in attempts:
        evaluation_map = attempt.get("evaluation") or {}
        if attempt.get("status") != "evaluated" or not evaluation_map:
            continue
        evaluated_attempts += 1
        paper = get_paper_by_id(attempt["paper_id"])
        if not paper:
            continue
        questions = {str(question["id"]): question for question in paper.get("questions", [])}
        for question_id, evaluation in evaluation_map.items():
            question = questions.get(str(question_id))
            if not question:
                continue
            max_marks = float(evaluation.get("max_marks") or question.get("marks") or 1)
            score = float(evaluation.get("score") or 0)
            concept = question.get("topic") or question.get("unit") or attempt.get("subject_name") or "Uncategorised"
            item = concepts.setdefault(concept, {
                "concept": concept,
                "attempts": 0,
                "students": set(),
                "score": 0.0,
                "max_score": 0.0,
                "correct": 0,
                "incorrect": 0,
            })
            item["attempts"] += 1
            item["students"].add(attempt["student_email"])
            item["score"] += score
            item["max_score"] += max_marks
            if score >= max_marks * 0.6:
                item["correct"] += 1
            else:
                item["incorrect"] += 1

    concept_rows = []
    for item in concepts.values():
        percentage = round((item["score"] / item["max_score"]) * 100) if item["max_score"] else 0
        concept_rows.append({
            "concept": item["concept"],
            "attempts": item["attempts"],
            "students": len(item["students"]),
            "accuracy": percentage,
            "correct": item["correct"],
            "incorrect": item["incorrect"],
            "status": "strong" if percentage >= 70 else "developing" if percentage >= 50 else "needs_attention",
        })
    concept_rows.sort(key=lambda item: (item["accuracy"], item["attempts"]), reverse=True)
    return jsonify({
        "evaluated_attempts": evaluated_attempts,
        "concepts": concept_rows,
        "strong_concepts": [item for item in concept_rows if item["status"] == "strong"],
        "developing_concepts": [item for item in concept_rows if item["status"] == "developing"],
        "needs_attention": sorted(
            [item for item in concept_rows if item["status"] == "needs_attention"],
            key=lambda item: item["accuracy"],
        ),
    })


@app.route("/api/teacher/attempts/<attempt_id>/grade", methods=["POST"])
def grade_teacher_attempt(attempt_id: str):
    """Automatically grade every submitted answer, with optional teacher corrections."""
    data = request.get_json() or {}
    conn = get_connection()
    row = conn.execute("SELECT * FROM attempts WHERE id = ?", (attempt_id,)).fetchone()
    conn.close()
    if row is None:
        return jsonify({"error": "Attempt not found"}), 404

    attempt = dict(row)
    paper = get_paper_by_id(attempt["paper_id"])
    if paper is None:
        return jsonify({"error": "Paper for this attempt was not found"}), 404

    answers = json.loads(attempt["answers"] or "{}")
    corrections = data.get("corrections") or {}
    evaluations = {}
    total_score = 0.0
    max_score = 0.0

    from utils.solution_generator import generate_smart_solution
    conn = get_connection()
    cursor = conn.cursor()
    try:
        for question in paper.get("questions", []):
            question_id = question["id"]
            max_marks = float(question.get("marks", 0))
            max_score += max_marks
            student_answer = answers.get(question_id, "")
            solution = generate_smart_solution(question, db_cursor=cursor, llm_gateway=llm_gateway)
            evaluation = evaluator.evaluate_answer(
                question_text=question.get("text", ""),
                model_answer=solution.get("solution", ""),
                student_answer=student_answer,
                max_marks=int(max_marks or 1),
            )
            correction = corrections.get(question_id)
            if isinstance(correction, dict):
                evaluation["score"] = max(0.0, min(max_marks, float(correction.get("score", evaluation["score"]))))
                if "feedback" in correction:
                    evaluation["feedback"] = str(correction["feedback"])
                evaluation["teacher_corrected"] = True
            total_score += float(evaluation.get("score", 0))
            evaluations[question_id] = evaluation
    finally:
        conn.close()

    updated = update_attempt_evaluation(attempt_id, evaluations, total_score, max_score)
    return jsonify(updated)


@app.route("/api/papers/<paper_id>", methods=["GET"])
def get_paper(paper_id):
    try:
        paper = get_paper_by_id(paper_id)
        if paper is None:
            return jsonify({"error": "Paper not found"}), 404
        return jsonify(paper)
    except Exception as e:
        logger.exception("Error getting paper")
        return jsonify({"error": str(e)}), 500


@app.route("/api/papers/<paper_id>", methods=["PUT"])
def modify_paper(paper_id):
    try:
        data = request.get_json()
        if not data:
            return jsonify({"error": "No JSON data provided"}), 400

        updated = update_paper(paper_id, data)
        if not updated:
            return jsonify({"error": "Paper not found or no changes made"}), 404
        return jsonify({"message": "Paper updated successfully", "paper": data})
    except Exception as e:
        logger.exception("Error updating paper")
        return jsonify({"error": str(e)}), 500


@app.route("/api/papers/<paper_id>/solutions", methods=["GET"])
def get_paper_solutions(paper_id):
    try:
        paper = get_paper_by_id(paper_id)
        if paper is None:
            return jsonify({"error": "Paper not found"}), 404

        questions = paper.get("questions", [])
        from utils.solution_generator import generate_smart_solution
        from database.db import get_connection
        conn = get_connection()
        cursor = conn.cursor()
        solutions = [generate_smart_solution(q, db_cursor=cursor, llm_gateway=llm_gateway) for q in questions]
        conn.close()
        return jsonify({
            "paper_id": paper_id,
            "subject": paper.get("subject", ""),
            "solutions": solutions
        })
    except Exception as e:
        logger.exception("Error generating solutions for paper")
        return jsonify({"error": str(e)}), 500


@app.route("/api/papers/<paper_id>", methods=["DELETE"])
def delete_paper(paper_id):
    try:
        deleted = delete_paper_by_id(paper_id)
        if not deleted:
            return jsonify({"error": "Paper not found"}), 404
        return jsonify({"message": "Paper deleted successfully"})
    except Exception as e:
        logger.exception("Error deleting paper")
        return jsonify({"error": str(e)}), 500


@app.route("/api/question-bank", methods=["GET"])
def get_question_bank():
    """Fetch filtered/paginated question bank entries."""
    try:
        subject = request.args.get("subject", None)
        difficulty = request.args.get("difficulty", None)
        question_type = request.args.get("type", None)
        search = request.args.get("search", None)
        page = int(request.args.get("page", 1))
        limit = int(request.args.get("limit", 20))

        result = get_pyq_questions_paginated(
            subject=subject,
            difficulty=difficulty,
            question_type=question_type,
            search=search,
            page=page,
            limit=limit,
        )
        return jsonify(result)
    except Exception as e:
        logger.exception("Error querying question bank")
        return jsonify({"error": str(e)}), 500


@app.route("/api/analytics", methods=["GET"])
def get_analytics():
    """Get overall platform analytics, topic frequencies, and PYQ stats."""
    try:
        data = get_all_pyq_analytics()
        return jsonify(data)
    except Exception as e:
        logger.exception("Error getting analytics data")
        return jsonify({"error": str(e)}), 500


@app.route("/api/evaluate-answer", methods=["POST"])
def evaluate_student_answer():
    """AI Auto-Grading & Semantic Evaluation of a student's answer."""
    try:
        data = request.get_json() or {}
        question = data.get("question", "").strip()
        model_answer = data.get("model_answer", "").strip()
        student_answer = data.get("student_answer", "").strip()
        max_marks = int(data.get("max_marks", 5))

        if not question:
            return jsonify({"error": "Question text is required"}), 400

        result = evaluator.evaluate_answer(
            question_text=question,
            model_answer=model_answer,
            student_answer=student_answer,
            max_marks=max_marks
        )
        return jsonify(result)
    except Exception as e:
        logger.exception("Error evaluating answer")
        return jsonify({"error": str(e)}), 500


@app.route("/api/generate-mcq", methods=["POST"])
def generate_mcqs():
    """Generate 4-option certification MCQs with distractors & explanations."""
    try:
        data = request.get_json() or {}
        subject = data.get("subject", "AWS").strip()
        topic = data.get("topic", "").strip()
        count = int(data.get("count", 5))

        mcqs = mcq_generator.generate_mcqs_for_subject(
            subject=subject,
            topic=topic,
            count=count,
            rag_engine=rag_engine
        )
        return jsonify({"subject": subject, "topic": topic, "mcqs": mcqs})
    except Exception as e:
        logger.exception("Error generating MCQs")
        return jsonify({"error": str(e)}), 500


@app.route("/api/papers/<paper_id>/pdf", methods=["GET"])
def export_pdf(paper_id):
    try:
        paper = get_paper_by_id(paper_id)
        if paper is None:
            return jsonify({"error": "Paper not found"}), 404

        filepath = pdf_gen.generate_pdf(paper)
        return send_file(
            filepath,
            as_attachment=True,
            download_name=f"{paper['subject'].replace(' ', '_')}_Question_Paper.pdf",
            mimetype="application/pdf",
        )
    except Exception as e:
        logger.exception("Error exporting PDF")
        return jsonify({"error": str(e)}), 500


@app.route("/api/papers/<paper_id>/export/moodle", methods=["GET"])
def export_moodle(paper_id):
    try:
        paper = get_paper_by_id(paper_id)
        if paper is None:
            return jsonify({"error": "Paper not found"}), 404
        xml_content = export_to_moodle_xml(paper)
        clean_name = paper.get("subject", "exam").replace(" ", "_")
        return Response(
            xml_content,
            mimetype="application/xml",
            headers={"Content-Disposition": f'attachment; filename="{clean_name}_moodle.xml"'}
        )
    except Exception as e:
        logger.exception("Error exporting Moodle XML")
        return jsonify({"error": str(e)}), 500


@app.route("/api/papers/<paper_id>/export/qti", methods=["GET"])
def export_qti(paper_id):
    try:
        paper = get_paper_by_id(paper_id)
        if paper is None:
            return jsonify({"error": "Paper not found"}), 404
        xml_content = export_to_qti_21(paper)
        clean_name = paper.get("subject", "exam").replace(" ", "_")
        return Response(
            xml_content,
            mimetype="application/xml",
            headers={"Content-Disposition": f'attachment; filename="{clean_name}_qti21.xml"'}
        )
    except Exception as e:
        logger.exception("Error exporting QTI 2.1")
        return jsonify({"error": str(e)}), 500


@app.route("/api/papers/<paper_id>/export/google-forms", methods=["GET"])
def export_google_forms_endpoint(paper_id):
    try:
        paper = get_paper_by_id(paper_id)
        if paper is None:
            return jsonify({"error": "Paper not found"}), 404
        forms_data = export_to_google_forms(paper)
        return jsonify(forms_data)
    except Exception as e:
        logger.exception("Error exporting Google Forms schema")
        return jsonify({"error": str(e)}), 500


@app.route("/api/papers/<paper_id>/export/docx", methods=["GET"])
def export_docx_endpoint(paper_id):
    try:
        paper = get_paper_by_id(paper_id)
        if paper is None:
            return jsonify({"error": "Paper not found"}), 404
        docx_stream = export_to_docx(paper)
        clean_name = paper.get("subject", "exam").replace(" ", "_")
        return send_file(
            docx_stream,
            as_attachment=True,
            download_name=f"{clean_name}_Question_Paper.docx",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    except Exception as e:
        logger.exception("Error exporting Word DOCX")
        return jsonify({"error": str(e)}), 500


@app.route("/api/subjects", methods=["GET"])
def list_subjects():
    subjects = [
        "AWS Cloud Fundamentals",
        "AWS Compute (EC2 & Auto Scaling)",
        "AWS Storage & Databases",
        "AWS Networking (VPC, Route 53, CloudFront)",
        "AWS Security & IAM",
        "AWS Serverless (Lambda, API Gateway, Step Functions)",
        "Docker & Containerization",
        "Kubernetes & Container Orchestration",
        "CI/CD Pipelines",
        "Jenkins",
        "Terraform & Infrastructure as Code",
        "Ansible & Configuration Management",
        "Linux Administration & Shell Scripting",
        "Git & Version Control",
        "Monitoring & Logging (CloudWatch, Prometheus, Grafana)",
        "Site Reliability Engineering (SRE)",
        "DevSecOps & Cloud Security",
        "Microservices Architecture",
    ]
    return jsonify(subjects)


@app.route("/api/analyze-syllabus", methods=["POST"])
def analyze_syllabus():
    try:
        data = request.get_json()
        syllabus = data.get("syllabus", "").strip()
        if not syllabus:
            return jsonify({"error": "Syllabus is required"}), 400

        topics = nlp.extract_topics(syllabus)
        units = nlp.extract_units(syllabus)

        return jsonify({"topics": topics, "units": units})
    except Exception as e:
        logger.exception("Error analyzing syllabus")
        return jsonify({"error": str(e)}), 500


@app.route("/api/upload-syllabus", methods=["POST"])
def upload_syllabus():
    """Extract syllabus text from uploaded PDF, DOCX, or text files."""
    try:
        if "file" not in request.files:
            return jsonify({"error": "No file uploaded"}), 400

        uploaded_file = request.files["file"]
        if not uploaded_file.filename:
            return jsonify({"error": "Empty filename"}), 400

        filename = uploaded_file.filename
        ext = os.path.splitext(filename)[1].lower()
        file_bytes = uploaded_file.read()

        if len(file_bytes) == 0:
            return jsonify({"error": "Uploaded file is empty"}), 400

        extracted_text = ""

        if ext == ".pdf":
            import io
            import pdfplumber
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                page_texts = []
                for page in pdf.pages:
                    txt = page.extract_text()
                    if txt:
                        page_texts.append(txt)
                extracted_text = "\n\n".join(page_texts)

        elif ext in [".docx", ".doc"]:
            import io
            try:
                import docx
                doc = docx.Document(io.BytesIO(file_bytes))
                lines = [p.text for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                        if row_text:
                            lines.append(row_text)
                extracted_text = "\n".join(lines)
            except Exception as docx_err:
                logger.warning("DOCX extraction error (%s). Trying raw text decode.", docx_err)
                extracted_text = file_bytes.decode("utf-8", errors="ignore")

        elif ext in [".txt", ".md", ".json"]:
            extracted_text = file_bytes.decode("utf-8", errors="ignore")
        else:
            extracted_text = file_bytes.decode("utf-8", errors="ignore")

        extracted_text = extracted_text.strip()
        if not extracted_text:
            return jsonify({"error": f"Could not extract readable text from '{filename}'"}), 400

        topics = nlp.extract_topics(extracted_text)
        units = nlp.extract_units(extracted_text)

        return jsonify({
            "filename": filename,
            "syllabus_text": extracted_text,
            "topics": topics,
            "units": units
        })

    except Exception as e:
        logger.exception("Error processing uploaded syllabus file")
        return jsonify({"error": str(e)}), 500



@app.route("/api/pyq-stats/<subject>", methods=["GET"])
def pyq_stats(subject):
    """Get PYQ statistics for a subject."""
    try:
        stats = get_pyq_stats_by_subject(subject)
        return jsonify(stats)
    except Exception as e:
        logger.exception("Error getting PYQ stats")
        return jsonify({"error": str(e)}), 500


@app.route("/api/health", methods=["GET"])
def health():
    return jsonify({
        "status": "ok",
        "message": "AI Question Paper Generator API is running",
        "llm": llm_gateway.get_status()
    })


@app.route("/api/llm/status", methods=["GET"])
def llm_status():
    """Get active LLM status and provider configuration."""
    return jsonify(llm_gateway.get_status())


@app.route("/api/db/backup", methods=["POST"])
def trigger_db_backup():
    """Trigger a live, consistent SQLite snapshot."""
    try:
        backup_path = backup_database()
        if backup_path and os.path.exists(backup_path):
            return jsonify({
                "status": "success",
                "message": "Database backup completed successfully",
                "backup_file": os.path.basename(backup_path),
                "backup_size_bytes": os.path.getsize(backup_path)
            }), 200
        return jsonify({"error": "Failed to create database snapshot"}), 500
    except Exception as e:
        logger.exception("Error executing database backup")
        return jsonify({"error": str(e)}), 500



def _classify_question_type(question: str) -> str:
    q_lower = question.lower()
    if any(kw in q_lower for kw in ["explain", "describe", "discuss", "elaborate"]):
        return "descriptive"
    elif any(kw in q_lower for kw in ["compare", "differentiate", "contrast"]):
        return "comparative"
    elif any(kw in q_lower for kw in ["derive", "prove", "calculate"]):
        return "analytical"
    elif any(kw in q_lower for kw in ["write a program", "implement", "code"]):
        return "programming"
    elif any(kw in q_lower for kw in ["diagram", "draw", "illustrate"]):
        return "diagrammatic"
    elif any(kw in q_lower for kw in ["list", "enumerate", "name"]):
        return "listing"
    elif any(kw in q_lower for kw in ["what is", "define", "short note"]):
        return "definition"
    else:
        return "descriptive"


if __name__ == "__main__":
    logger.info("Starting AI Question Paper Generator API...")
    logger.info("Server running at http://%s:%d", FLASK_HOST, FLASK_PORT)
    app.run(
        host=FLASK_HOST,
        port=FLASK_PORT,
        debug=FLASK_DEBUG,
        threaded=True,        # Handle multiple requests concurrently
        use_reloader=False,   # Prevent double model-loading in debug mode
    )
